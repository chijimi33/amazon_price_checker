import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import TableColumn

from catalog_excel import build_catalog_from_workbook, convert_value


REPOSITORY_ROOT = Path(__file__).parents[1]
WORKBOOK = REPOSITORY_ROOT / "catalog" / "quality_catalog.xlsx"


def table_geometry(sheet, table_name):
    table = sheet.tables[table_name]
    return table, range_boundaries(table.ref)


def write_first_empty_row(sheet, table_name, values):
    _, (min_col, min_row, max_col, max_row) = table_geometry(sheet, table_name)
    headers = {
        str(sheet.cell(min_row, column).value): column
        for column in range(min_col, max_col + 1)
    }
    for row in range(min_row + 1, max_row + 1):
        if all(sheet.cell(row, column).value in (None, "") for column in range(min_col, max_col + 1)):
            for key, value in values.items():
                sheet.cell(row, headers[key]).value = value
            return row
    raise AssertionError(f"{table_name} has no reserved row")


def fill_minimum_ssd(workbook):
    return write_first_empty_row(
        workbook["SSD"],
        "SSDCatalog",
        {
            "product_id": "test-ssd",
            "brand": "Example",
            "model": "EX-2TB",
            "display_name": "Example 2TB SSD",
            "status": "research_required",
            "tier": "unrated",
            "capacity_gb": 2000,
            "nand_type": "TLC",
        },
    )


class CatalogExcelTest(unittest.TestCase):
    def copy_workbook(self):
        temporary = tempfile.TemporaryDirectory()
        destination = Path(temporary.name) / "quality_catalog.xlsx"
        shutil.copy2(WORKBOOK, destination)
        return temporary, destination

    def test_repository_workbook_is_valid_and_matches_json(self):
        import json

        result = build_catalog_from_workbook(WORKBOOK)
        expected = json.loads(
            (REPOSITORY_ROOT / "catalog" / "quality_catalog.json").read_text(encoding="utf-8")
        )
        self.assertEqual(result.errors, [])
        self.assertEqual(result.warnings, [])
        self.assertEqual(result.catalog, expected)

    def test_cpu_catalog_contains_supported_generations_and_performance_data(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        products = {
            product["id"]: product
            for product in result.catalog["products"]
            if product["category"] == "cpu"
        }
        self.assertGreaterEqual(len(products), 112)

        ryzen = products["amd-ryzen-5-5600x"]
        self.assertEqual(ryzen["specs"]["generation"], "Ryzen 5000")
        self.assertEqual(ryzen["specs"]["release_date"], "2020-11-05")
        self.assertGreater(ryzen["specs"]["passmark_cpu_mark"], 0)
        self.assertGreater(ryzen["specs"]["passmark_single_thread_mark"], 0)

        intel = products["intel-core-i5-12400f"]
        self.assertEqual(intel["specs"]["generation"], "12th Gen Core")
        self.assertEqual(intel["specs"]["release_date"], "2022-01-04")
        self.assertFalse(intel["specs"]["integrated_graphics"])
        self.assertGreaterEqual(len(intel["quality"]["evidence"]), 2)

    def test_gpu_chip_catalog_contains_supported_series_and_benchmarks(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        chips = [
            product
            for product in result.catalog["products"]
            if product["category"] == "gpu_chip"
        ]
        self.assertGreaterEqual(len(chips), 64)
        self.assertGreaterEqual(sum(chip["brand"] == "NVIDIA" for chip in chips), 38)
        self.assertGreaterEqual(sum(chip["brand"] == "AMD" for chip in chips), 26)

        products = {product["id"]: product for product in chips}
        rtx = products["nvidia-geforce-rtx-5090"]
        self.assertEqual(rtx["specs"]["generation"], "GeForce RTX 50 Series")
        self.assertEqual(rtx["specs"]["vram_gb"], 32)
        self.assertGreater(rtx["specs"]["passmark_g3d_mark"], 0)

        radeon = products["amd-radeon-rx-9070-gre"]
        self.assertEqual(radeon["specs"]["generation"], "Radeon RX 9000 Series")
        self.assertEqual(radeon["specs"]["release_date"], "2026-06-02")
        self.assertGreaterEqual(len(radeon["quality"]["evidence"]), 2)

        rx6500 = products["amd-radeon-rx-6500-xt-8gb"]
        self.assertEqual(rx6500["specs"]["vram_gb"], 8)
        self.assertGreaterEqual(len(rx6500["quality"]["evidence"]), 2)

        rx6700 = products["amd-radeon-rx-6700"]
        self.assertEqual(rx6700["specs"]["release_date"], "2021-06-09")
        self.assertEqual(rx6700["specs"]["release_date_precision"], "exact")

        profile = next(
            profile
            for profile in result.catalog["profiles"]
            if profile["id"] == "gpu-1440p-16gb"
        )
        cooler_rule = next(
            rule
            for rule in profile["preferences"]
            if rule["path"] == "specs.cooler_quality_tier"
        )
        self.assertIn("S", cooler_rule["value"])

    def test_gpu_board_data_validations_match_their_columns(self):
        workbook = load_workbook(WORKBOOK, data_only=False)
        validations = {
            (str(validation.sqref), validation.formula1)
            for validation in workbook["GPU"].data_validations.dataValidation
        }

        self.assertIn(("W5:X104", '"true,false"'), validations)
        self.assertIn(("AA5:AA104", "'Lists'!$P$5:$P$8"), validations)
        self.assertIn(("AD5:AD104", "'Lists'!$R$5:$R$13"), validations)
        self.assertIn(("AE5:AE104", "'Lists'!$Q$5:$Q$9"), validations)
        self.assertNotIn(("V5:V104", "'Lists'!$P$5:$P$8"), validations)
        self.assertNotIn(("X5:X104", "'Lists'!$R$5:$R$13"), validations)
        self.assertNotIn(("Y5:Y104", "'Lists'!$Q$5:$Q$9"), validations)

    def test_expanded_tables_keep_input_rules_and_status_formatting(self):
        workbook = load_workbook(WORKBOOK, data_only=False)
        expected_validations = {
            "CPU": {
                ("J5:J136", "'Lists'!$A$5:$A$9", False),
                ("K5:K136", "'Lists'!$B$5:$B$9", False),
                ("R5:R136", "'Lists'!$N$5:$N$8", False),
                ("X5:Z136", "'Lists'!$C$5:$C$6", False),
                ("AA5:AA136", "'Lists'!$O$5:$O$8", False),
                ("AF5:AF136", "'Lists'!$AM$5:$AM$8", False),
            },
            "SSD": {
                ("J5:J701", "'Lists'!$A$5:$A$9", False),
                ("K5:K701", "'Lists'!$B$5:$B$9", False),
                ("S5:S701", "'Lists'!$U$5:$U$10", False),
                ("U5:U701", "'Lists'!$V$5:$V$9", False),
                ("V5:V701", "'Lists'!$C$5:$C$6", False),
                ("W5:W701", "'Lists'!$W$5:$W$9", False),
                ("Z5:Z701", "'Lists'!$X$5:$X$8", False),
                ("AD5:AD701", "'Lists'!$C$5:$C$6", False),
            },
            "PSU": {
                ("J5:J485", "'Lists'!$A$5:$A$9", False),
                ("K5:K485", "'Lists'!$B$5:$B$9", False),
                ("S5:S485", "'Lists'!$Y$5:$Y$8", False),
                ("T5:T485", "'Lists'!$C$5:$C$6", False),
                ("V5:V485", "'Lists'!$C$5:$C$6", False),
                ("W5:W485", "'Lists'!$Z$5:$Z$9", False),
                ("X5:X485", "'Lists'!$AA$5:$AA$10", False),
                ("Y5:Y485", "'Lists'!$AB$5:$AB$9", False),
                ("Z5:Z485", "'Lists'!$AC$5:$AC$7", False),
                ("AB5:AB485", "'Lists'!$AD$5:$AD$7", False),
            },
            "Motherboard": {
                ("J5:J487", "'Lists'!$A$5:$A$9", False),
                ("K5:K487", "'Lists'!$B$5:$B$9", False),
                ("T5:T487", "'Lists'!$AM$5:$AM$8", True),
                ("V5:V487", "'Lists'!$N$5:$N$8", True),
                ("X5:X487", "'Lists'!$S$5:$S$6", True),
                ("Y5:Y487", "'Lists'!$AE$5:$AE$8", True),
                ("AF5:AF487", "'Lists'!$AS$5:$AS$7", True),
                ("AR5:AR487", "'Lists'!$AG$5:$AG$9", True),
                ("AU5:AU487", "'Lists'!$AR$5:$AR$7", True),
                ("BG5:BI487", "'Lists'!$C$5:$C$6", True),
            },
            "MotherboardSlots": {
                ("B5:B2678", "'Motherboard'!$A$5:$A$487", False),
                ("D5:D2678", "'Lists'!$AN$5:$AN$6", False),
                ("G5:G2678", "'Lists'!$AO$5:$AO$7", True),
                ("I5:I2678", "'Lists'!$C$5:$C$6", True),
            },
            "MotherboardUSB": {
                ("B5:B1979", "'Motherboard'!$A$5:$A$487", False),
                ("C5:C1979", "'Lists'!$AP$5:$AP$6", False),
                ("F5:F1979", "'Lists'!$AQ$5:$AQ$6", False),
            },
            "Evidence": {
                ("C5:C2541", "'Lists'!$D$5:$D$10", False),
            },
            "Monitor": {
                ("J5:J99", "'Lists'!$A$5:$A$9", False),
                ("K5:K99", "'Lists'!$B$5:$B$9", False),
                ("S5:S99", "'Lists'!$AH$5:$AH$8", False),
                ("U5:U99", "'Lists'!$C$5:$C$6", False),
                ("V5:V99", "'Lists'!$C$5:$C$6", False),
                ("X5:X99", "'Lists'!$AI$5:$AI$9", False),
                ("Z5:Z99", "'Lists'!$C$5:$C$6", False),
                ("AB5:AB99", "'Lists'!$AJ$5:$AJ$12", False),
            },
            "FieldDefinitions": {
                ("D5:D317", "'Lists'!$G$5:$G$11", False),
                ("E5:E317", "'Lists'!$C$5:$C$6", False),
                ("I5:I317", "'Lists'!$C$5:$C$6", False),
            },
        }

        for sheet_name, expected in expected_validations.items():
            with self.subTest(sheet=sheet_name):
                actual = {
                    (
                        str(validation.sqref),
                        validation.formula1,
                        bool(validation.allow_blank),
                    )
                    for validation in workbook[
                        sheet_name
                    ].data_validations.dataValidation
                }
                self.assertEqual(actual, expected)

        expected_status_ranges = {
            "CPU": "J5:J136",
            "GPU": "J5:J104",
            "GPUChips": "F5:F87",
            "Memory": "J5:J24",
            "SSD": "J5:J701",
            "PSU": "J5:J485",
            "Motherboard": "J5:J487",
            "Monitor": "J5:J99",
        }
        for sheet_name, expected_range in expected_status_ranges.items():
            with self.subTest(status_sheet=sheet_name):
                ranges = {
                    str(item.sqref)
                    for item in workbook[sheet_name].conditional_formatting
                }
                self.assertIn(expected_range, ranges)
                if expected_range not in {"J5:J24", "F5:F24"}:
                    self.assertNotIn("J5:J24", ranges)

        self.assertEqual(workbook["Lists"]["D10"].value, "price_com")

    def test_gpu_chip_sheet_is_visible_and_adjacent_to_gpu(self):
        workbook = load_workbook(WORKBOOK, data_only=False)
        self.assertEqual(workbook["GPUChips"].sheet_state, "visible")
        self.assertEqual(
            workbook.sheetnames.index("GPUChips"),
            workbook.sheetnames.index("GPU") + 1,
        )

    def test_gpu_chip_sheet_has_formula_driven_performance_view(self):
        workbook = load_workbook(WORKBOOK, data_only=False)
        sheet = workbook["GPUChips"]
        self.assertEqual(
            [sheet.cell(4, column).value for column in range(26, 30)],
            [
                "performance_index_rtx4060_100",
                "performance_rank",
                "performance_class",
                "passmark_g3d_per_watt",
            ],
        )
        self.assertEqual(sheet["AA3"].value, "nvidia-geforce-rtx-4060")
        self.assertIn("$AA$3", sheet["Z5"].value)
        self.assertIn("COUNTIF", sheet["AA5"].value)
        self.assertIn("'Lists'!$AL$9", sheet["AB5"].value)
        self.assertIn("V5/X5", sheet["AC5"].value)
        self.assertEqual(workbook["Lists"]["AK5"].value, "エントリー")
        self.assertEqual(workbook["Lists"]["AL9"].value, 170)

    def test_gpu_board_catalog_contains_initial_exact_skus(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        products = {
            product["id"]: product
            for product in result.catalog["products"]
            if product["category"] == "gpu"
        }
        self.assertGreaterEqual(len(products), 44)
        self.assertTrue(
            {"ASUS", "MSI", "GIGABYTE", "SAPPHIRE", "PowerColor", "ASRock"}
            .issubset({product["brand"] for product in products.values()})
        )

        generations = {
            product["specs"]["gpu_chip"]["generation"]
            for product in products.values()
        }
        self.assertTrue(
            {
                "GeForce RTX 20 Series",
                "GeForce RTX 30 Series",
                "GeForce RTX 40 Series",
                "GeForce RTX 50 Series",
                "Radeon RX 6000 Series",
                "Radeon RX 7000 Series",
                "Radeon RX 9000 Series",
            }.issubset(generations)
        )

        asus = products["gpu-asus-dual-rtx3060-o12g-v2"]
        self.assertEqual(
            asus["identifiers"]["part_numbers"],
            ["DUAL-RTX3060-O12G-V2"],
        )
        self.assertEqual(asus["specs"]["revision"], "V2")
        self.assertEqual(asus["specs"]["length_mm"], 200)

        sapphire = products["gpu-sapphire-11348-03-20g"]
        self.assertEqual(sapphire["specs"]["fan_count"], 3)
        self.assertEqual(sapphire["specs"]["slot_width"], 3)
        self.assertEqual(sapphire["specs"]["power_connector_standard"], "2x8pin")
        self.assertEqual(sapphire["specs"]["tdp_w"], 304)

        msi = products["gpu-msi-g5070-12gtc"]
        self.assertEqual(msi["identifiers"]["part_numbers"], ["G5070-12GTC"])
        self.assertEqual(msi["specs"]["gpu_chip"]["vram_gb"], 12)
        self.assertEqual(msi["specs"]["power_connector_standard"], "16pin")

        asrock = products["gpu-asrock-rx6400-cli-4g"]
        self.assertEqual(asrock["specs"]["length_mm"], 162)
        self.assertEqual(asrock["specs"]["power_connector_standard"], "none")
        self.assertEqual(
            asrock["quality"]["strengths"],
            ["162mm長", "シングルファン", "補助電源不要"],
        )

        for product in products.values():
            self.assertTrue(
                any(
                    evidence["kind"] == "manufacturer"
                    for evidence in product["quality"]["evidence"]
                )
            )

    def test_gpu_board_lifecycle_can_be_human_approved(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            sheet = workbook["GPU"]
            headers = {
                str(sheet.cell(4, column).value): column
                for column in range(1, sheet.max_column + 1)
            }
            row = next(
                row
                for row in range(5, sheet.max_row + 1)
                if sheet.cell(row, headers["product_id"]).value
                == "gpu-asus-dual-rtx3060-o12g-v2"
            )
            sheet.cell(row, headers["status"]).value = "approved"
            sheet.cell(row, headers["tier"]).value = "B"
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertEqual(result.errors, [])
            product = next(
                product
                for product in result.catalog["products"]
                if product["id"] == "gpu-asus-dual-rtx3060-o12g-v2"
            )
            self.assertEqual(product["quality"]["status"], "approved")
            self.assertEqual(product["quality"]["tier"], "B")
        finally:
            temporary.cleanup()

    def test_gpu_board_resolves_normalized_chip_specs(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            write_first_empty_row(
                workbook["GPUChips"],
                "GPUChipsCatalog",
                {
                    "product_id": "test-gpu-chip",
                    "brand": "NVIDIA",
                    "model": "Test GPU 16GB",
                    "display_name": "NVIDIA Test GPU 16GB",
                    "status": "research_required",
                    "tier": "unrated",
                    "vram_gb": 16,
                    "passmark_g3d_mark": 30000,
                },
            )
            write_first_empty_row(
                workbook["GPU"],
                "GPUCatalog",
                {
                    "product_id": "test-gpu-board",
                    "brand": "ExampleBoard",
                    "model": "EX-GPU-16",
                    "display_name": "ExampleBoard EX-GPU-16",
                    "status": "research_required",
                    "tier": "unrated",
                    "gpu_chip_id": "test-gpu-chip",
                    "board_series": "Example Series",
                    "domestic_warranty_years": 2,
                },
            )
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertEqual(result.errors, [])
            board = next(
                product
                for product in result.catalog["products"]
                if product["id"] == "test-gpu-board"
            )
            self.assertEqual(board["specs"]["gpu_chip"]["id"], "test-gpu-chip")
            self.assertEqual(board["specs"]["gpu_chip"]["vram_gb"], 16)
            self.assertEqual(board["specs"]["gpu_chip"]["passmark_g3d_mark"], 30000)
        finally:
            temporary.cleanup()

    def test_memory_catalog_contains_initial_exact_sku_kits(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        products = {
            product["id"]: product
            for product in result.catalog["products"]
            if product["category"] == "memory"
        }
        self.assertGreaterEqual(len(products), 16)
        self.assertEqual(
            {product["brand"] for product in products.values()},
            {"Corsair", "G.SKILL", "Kingston", "Crucial"},
        )

        corsair = products["memory-corsair-cmk32gx5m2b6000z30"]
        self.assertEqual(
            corsair["identifiers"]["part_numbers"],
            ["CMK32GX5M2B6000Z30"],
        )
        self.assertEqual(corsair["specs"]["memory_type"], "DDR5")
        self.assertEqual(corsair["specs"]["total_capacity_gb"], 32)
        self.assertEqual(corsair["specs"]["module_count"], 2)
        self.assertEqual(corsair["specs"]["data_rate_mt_s"], 6000)
        self.assertEqual(corsair["specs"]["default_data_rate_mt_s"], 4800)
        self.assertTrue(corsair["specs"]["expo"])
        self.assertEqual(corsair["specs"]["cas_latency"], 30)

        kingston = products["memory-kingston-kf560c30bbek2-64"]
        self.assertEqual(kingston["specs"]["module_height_mm"], 34.9)
        self.assertEqual(kingston["specs"]["warranty_class"], "limited_lifetime")

        crucial = products["memory-crucial-cp2k16g4dfra32a"]
        self.assertEqual(crucial["specs"]["memory_type"], "DDR4")
        self.assertFalse(crucial["specs"]["expo"])
        self.assertEqual(crucial["specs"]["voltage_v"], 1.2)

        initial_ids = {
            "memory-corsair-cmk32gx5m2b6000z30",
            "memory-corsair-cmk64gx5m2b6000z30",
            "memory-corsair-cmk32gx4m2e3200c16",
            "memory-corsair-cmk64gx4m2e3200c16",
            "memory-gskill-f5-6000j3038f16gx2-fx5",
            "memory-gskill-f5-6000j3040g32gx2-fx5",
            "memory-gskill-f4-3600c16d-32gvkc",
            "memory-gskill-f4-3600c18d-64gvk",
            "memory-kingston-kf560c30bbek2-32",
            "memory-kingston-kf560c30bbek2-64",
            "memory-kingston-kf432c16bbk2-32",
            "memory-kingston-kf432c16bbk2-64",
            "memory-crucial-cp2k16g60c36u5b",
            "memory-crucial-cp2k32g60c40u5b",
            "memory-crucial-cp2k16g4dfra32a",
            "memory-crucial-cp2k32g4dfra32a",
        }
        expected_default_data_rates = {
            "memory-corsair-cmk32gx5m2b6000z30": 4800,
            "memory-corsair-cmk64gx5m2b6000z30": 4800,
            "memory-corsair-cmk32gx4m2e3200c16": 2133,
            "memory-corsair-cmk64gx4m2e3200c16": 2133,
            "memory-gskill-f5-6000j3038f16gx2-fx5": 4800,
            "memory-gskill-f5-6000j3040g32gx2-fx5": 4800,
            "memory-gskill-f4-3600c16d-32gvkc": 2133,
            "memory-gskill-f4-3600c18d-64gvk": 2666,
            "memory-kingston-kf560c30bbek2-32": 4800,
            "memory-kingston-kf560c30bbek2-64": 4800,
            "memory-kingston-kf432c16bbk2-32": 2400,
            "memory-kingston-kf432c16bbk2-64": 2400,
            "memory-crucial-cp2k16g60c36u5b": 5600,
            "memory-crucial-cp2k32g60c40u5b": 5600,
            "memory-crucial-cp2k16g4dfra32a": 3200,
            "memory-crucial-cp2k32g4dfra32a": 3200,
        }
        self.assertTrue(initial_ids.issubset(products))
        for product_id in initial_ids:
            product = products[product_id]
            self.assertEqual(
                product["specs"]["default_data_rate_mt_s"],
                expected_default_data_rates[product_id],
            )
            self.assertEqual(product["quality"]["status"], "research_required")
            self.assertEqual(product["quality"]["tier"], "unrated")
            self.assertGreaterEqual(len(product["quality"]["evidence"]), 1)
            self.assertTrue(
                any(
                    item["kind"] == "manufacturer"
                    for item in product["quality"]["evidence"]
                )
            )

    def test_monitor_catalog_contains_chimolog_a_or_better_products(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        monitors = [
            product
            for product in result.catalog["products"]
            if product["category"] == "monitor"
        ]
        self.assertEqual(len(monitors), 75)
        self.assertEqual(sum(product["quality"]["tier"] == "S" for product in monitors), 9)
        self.assertEqual(sum(product["quality"]["tier"] == "A" for product in monitors), 66)
        self.assertEqual(
            sum(product["quality"]["status"] == "discontinued" for product in monitors),
            4,
        )
        self.assertTrue(
            all(len(product["quality"]["evidence"]) >= 2 for product in monitors)
        )

        products = {product["id"]: product for product in monitors}
        p275ms_plus = products["titan-army-p275ms-plus"]
        self.assertEqual(p275ms_plus["specs"]["resolution"], "2560x1440")
        self.assertEqual(p275ms_plus["specs"]["max_refresh_hz"], 320)
        self.assertTrue(p275ms_plus["specs"]["measured_response_reviewed"])

        model_conflict = products["msi-g274qpx"]
        self.assertEqual(
            model_conflict["quality"]["status"],
            "research_required",
        )
        self.assertIn("型番が不一致", model_conflict["quality"]["risk_summary"])

    def test_ssd_and_psu_candidates_follow_selected_scope(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        ssd = [
            product
            for product in result.catalog["products"]
            if product["category"] == "ssd"
        ]
        psu = [
            product
            for product in result.catalog["products"]
            if product["category"] == "psu"
        ]
        self.assertEqual(len(ssd), 677)
        self.assertEqual(len(psu), 461)

        strict_ssd_brands = {
            "Hanye",
            "WINTEN",
            "SPD",
            "AGI",
            "KOWIN",
            "addlink",
            "JNH",
            "HI-DISC",
            "Verbatim",
            "IODATA",
            "エレコム",
            "ロジテック",
        }
        excluded_psu_brands = {
            "ADATA",
            "Lian Li",
            "IN WIN",
            "In Win",
            "Enhance",
            "Segotep",
            "PCCOOLER",
            "ZALMAN",
            "Sharkoon",
            "darkFlash",
        }
        self.assertFalse({product["brand"] for product in ssd} & strict_ssd_brands)
        self.assertFalse({product["brand"] for product in psu} & excluded_psu_brands)
        self.assertIn("ドスパラセレクト", {product["brand"] for product in psu})

        for product in ssd + psu:
            self.assertEqual(product["quality"]["status"], "research_required")
            self.assertEqual(product["quality"]["tier"], "unrated")
            self.assertTrue(product["identifiers"].get("kakaku_ids"))
            self.assertTrue(
                any(
                    item["kind"] == "technical_database"
                    and item["url"].startswith("https://kakaku.com/item/")
                    for item in product["quality"]["evidence"]
                )
            )

        workbook = load_workbook(WORKBOOK, data_only=False, read_only=False)
        self.assertEqual(workbook["Search"]["B9"].value, "=COUNTA('SSD'!$A$5:$A$5000)")
        self.assertEqual(workbook["Search"]["B10"].value, "=COUNTA('PSU'!$A$5:$A$5000)")

    def test_motherboard_catalog_contains_normalized_slots_and_usb(self):
        result = build_catalog_from_workbook(WORKBOOK)
        self.assertEqual(result.errors, [])
        products = {
            product["id"]: product
            for product in result.catalog["products"]
            if product["category"] == "motherboard"
        }
        self.assertEqual(len(products), 463)
        self.assertEqual(
            {
                brand: sum(product["brand"] == brand for product in products.values())
                for brand in ("ASRock", "ASUS", "GIGABYTE", "MSI")
            },
            {"ASRock": 143, "ASUS": 105, "GIGABYTE": 112, "MSI": 103},
        )
        self.assertEqual(
            {
                socket: sum(
                    product["specs"]["socket"] == socket
                    for product in products.values()
                )
                for socket in ("AM4", "AM5", "LGA1700", "LGA1851")
            },
            {"AM4": 46, "AM5": 220, "LGA1700": 86, "LGA1851": 111},
        )
        self.assertTrue(
            all(product["identifiers"].get("kakaku_ids") for product in products.values())
        )
        self.assertTrue(
            all(
                any(item["kind"] == "price_com" for item in product["quality"]["evidence"])
                for product in products.values()
            )
        )
        self.assertEqual(
            sum(
                any(
                    item["kind"] == "manufacturer"
                    for item in product["quality"]["evidence"]
                )
                for product in products.values()
            ),
            351,
        )

        msi_b650 = products["motherboard-msi-mag-b650-tomahawk-wifi"]
        self.assertEqual(msi_b650["specs"]["m2_slots"], 3)
        self.assertEqual(msi_b650["specs"]["m2_heatsink_slots"], 2)
        self.assertEqual(msi_b650["specs"]["rear_usb_total_count"], 10)
        self.assertEqual(len(msi_b650["specs"]["slots"]), 6)
        self.assertEqual(len(msi_b650["specs"]["usb_ports"]), 5)
        self.assertIn("lane_sharing", msi_b650["quality"]["risk_flags"])

        asrock_z890 = products["motherboard-asrock-z890-steel-legend-wifi"]
        self.assertEqual(asrock_z890["specs"]["rear_usb_fastest_gbps"], 40)
        self.assertEqual(asrock_z890["specs"]["rear_usb4_type_c_count"], 2)
        self.assertEqual(asrock_z890["specs"]["pump_capable_headers"], 7)
        self.assertEqual(asrock_z890["specs"]["m2_heatsink_slots"], 3)

        self.assertEqual(
            sum(bool(product["specs"].get("slots")) for product in products.values()),
            453,
        )
        self.assertEqual(
            sum(bool(product["specs"].get("usb_ports")) for product in products.values()),
            453,
        )
        self.assertEqual(
            sum(len(product["specs"].get("slots", [])) for product in products.values()),
            2634,
        )
        self.assertEqual(
            sum(
                len(product["specs"].get("usb_ports", []))
                for product in products.values()
            ),
            1925,
        )

        for product in products.values():
            m2_slots = [
                slot
                for slot in product["specs"].get("slots", [])
                if slot["type"] == "m2_storage"
            ]
            if not m2_slots:
                continue
            self.assertEqual(product["specs"]["m2_slots"], len(m2_slots))
            if all(isinstance(slot.get("interface_generation"), int) for slot in m2_slots):
                self.assertEqual(
                    product["specs"]["pcie5_m2_slots"],
                    sum(slot["interface_generation"] == 5 for slot in m2_slots),
                )
            if all(isinstance(slot.get("heatsink"), bool) for slot in m2_slots):
                self.assertEqual(
                    product["specs"]["m2_heatsink_slots"],
                    sum(slot["heatsink"] for slot in m2_slots),
                )

    def test_motherboard_child_sheets_are_visible_and_adjacent(self):
        workbook = load_workbook(WORKBOOK, data_only=False, read_only=False)
        names = workbook.sheetnames
        motherboard_index = names.index("Motherboard")
        self.assertEqual(names[motherboard_index + 1], "MotherboardSlots")
        self.assertEqual(names[motherboard_index + 2], "MotherboardUSB")
        self.assertEqual(workbook["MotherboardSlots"].sheet_state, "visible")
        self.assertEqual(workbook["MotherboardUSB"].sheet_state, "visible")
        _, slot_geometry = table_geometry(
            workbook["MotherboardSlots"], "MotherboardSlotsCatalog"
        )
        self.assertEqual(
            slot_geometry[:3],
            (1, 4, 13),
        )
        slot_rows = sum(
            workbook["MotherboardSlots"].cell(row, 1).value not in (None, "")
            for row in range(slot_geometry[1] + 1, slot_geometry[3] + 1)
        )
        self.assertEqual(slot_rows, 2634)
        self.assertGreaterEqual(slot_geometry[3] - slot_geometry[1], slot_rows + 40)

        _, usb_geometry = table_geometry(
            workbook["MotherboardUSB"], "MotherboardUSBCatalog"
        )
        self.assertEqual(
            usb_geometry[:3],
            (1, 4, 10),
        )
        usb_rows = sum(
            workbook["MotherboardUSB"].cell(row, 1).value not in (None, "")
            for row in range(usb_geometry[1] + 1, usb_geometry[3] + 1)
        )
        self.assertEqual(usb_rows, 1925)
        self.assertGreaterEqual(usb_geometry[3] - usb_geometry[1], usb_rows + 50)

    def test_motherboard_usb_parent_aggregate_mismatch_is_rejected(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            sheet = workbook["Motherboard"]
            _, (min_col, min_row, max_col, _) = table_geometry(
                sheet,
                "MotherboardCatalog",
            )
            headers = {
                str(sheet.cell(min_row, column).value): column
                for column in range(min_col, max_col + 1)
            }
            total_column = headers["rear_usb_total_count"]
            sheet.cell(min_row + 1, total_column).value += 1
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertTrue(
                any(
                    "MotherboardUSB集計" in error
                    and "rear_usb_total_count" in error
                    for error in result.errors
                )
            )
        finally:
            temporary.cleanup()

    def test_motherboard_front_usb_parent_aggregate_mismatch_is_rejected(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            sheet = workbook["MotherboardUSB"]
            _, (min_col, min_row, max_col, max_row) = table_geometry(
                sheet,
                "MotherboardUSBCatalog",
            )
            headers = {
                str(sheet.cell(min_row, column).value): column
                for column in range(min_col, max_col + 1)
            }
            for row in range(min_row + 1, max_row + 1):
                if (
                    sheet.cell(row, headers["product_id"]).value
                    == "motherboard-msi-mag-b650-tomahawk-wifi"
                    and sheet.cell(row, headers["location"]).value == "front_header"
                    and sheet.cell(row, headers["connector_type"]).value == "Type-C"
                ):
                    sheet.cell(row, headers["usb_max_speed_gbps"]).value = 20
                    break
            else:
                self.fail("MSI B650のフロントUSB Type-C行が見つかりません")
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertTrue(
                any(
                    "MotherboardUSB集計" in error
                    and "front_usb_c_max_speed_gbps" in error
                    for error in result.errors
                )
            )
        finally:
            temporary.cleanup()

    def test_motherboard_m2_parent_aggregate_mismatches_are_rejected(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            sheet = workbook["Motherboard"]
            _, (min_col, min_row, max_col, _) = table_geometry(
                sheet,
                "MotherboardCatalog",
            )
            headers = {
                str(sheet.cell(min_row, column).value): column
                for column in range(min_col, max_col + 1)
            }
            product_id_column = headers["product_id"]
            target_row = next(
                row
                for row in range(min_row + 1, sheet.max_row + 1)
                if sheet.cell(row, product_id_column).value
                == "motherboard-asrock-z890-steel-legend-wifi"
            )
            for key in ("m2_slots", "pcie5_m2_slots", "m2_heatsink_slots"):
                cell = sheet.cell(target_row, headers[key])
                cell.value += 1
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            for key in ("m2_slots", "pcie5_m2_slots", "m2_heatsink_slots"):
                with self.subTest(key=key):
                    self.assertTrue(
                        any(
                            "MotherboardSlots集計" in error and key in error
                            for error in result.errors
                        )
                    )
        finally:
            temporary.cleanup()

    def test_motherboard_child_enums_are_rejected_when_unknown(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            slots = workbook["MotherboardSlots"]
            _, (slot_min_col, slot_header_row, slot_max_col, _) = table_geometry(
                slots,
                "MotherboardSlotsCatalog",
            )
            slot_headers = {
                str(slots.cell(slot_header_row, column).value): column
                for column in range(slot_min_col, slot_max_col + 1)
            }
            slots.cell(slot_header_row + 1, slot_headers["slot_type"]).value = "unknown"

            usb = workbook["MotherboardUSB"]
            _, (usb_min_col, usb_header_row, usb_max_col, _) = table_geometry(
                usb,
                "MotherboardUSBCatalog",
            )
            usb_headers = {
                str(usb.cell(usb_header_row, column).value): column
                for column in range(usb_min_col, usb_max_col + 1)
            }
            usb.cell(usb_header_row + 1, usb_headers["connector_type"]).value = "Type-X"
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertTrue(any("slot_type: 未対応値" in error for error in result.errors))
            self.assertTrue(
                any("connector_type: 未対応値" in error for error in result.errors)
            )
        finally:
            temporary.cleanup()

    def test_new_spec_column_is_mapped_without_converter_change(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            target_row = fill_minimum_ssd(workbook)
            sheet = workbook["SSD"]
            table, (min_col, min_row, max_col, max_row) = table_geometry(sheet, "SSDCatalog")
            new_column = max_col + 1
            sheet.cell(min_row, new_column).value = "test_metric"
            sheet.cell(target_row, new_column).value = 12.5
            table.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(new_column)}{max_row}"
            )
            table.tableColumns.append(
                TableColumn(id=len(table.tableColumns) + 1, name="test_metric")
            )
            write_first_empty_row(
                workbook["FieldDefinitions"],
                "FieldDefinitionsCatalog",
                {
                    "sheet_name": "SSD",
                    "column_name": "test_metric",
                    "json_path": "specs.test_metric",
                    "data_type": "number",
                    "required": False,
                    "description": "拡張テスト",
                    "active": True,
                },
            )
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertEqual(result.errors, [])
            product = next(
                product
                for product in result.catalog["products"]
                if product["id"] == "test-ssd"
            )
            self.assertEqual(product["specs"]["test_metric"], 12.5)
        finally:
            temporary.cleanup()

    def test_duplicate_json_path_is_rejected_before_silent_overwrite(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            sheet = workbook["FieldDefinitions"]
            _, (min_col, min_row, max_col, max_row) = table_geometry(
                sheet,
                "FieldDefinitionsCatalog",
            )
            headers = {
                str(sheet.cell(min_row, column).value): column
                for column in range(min_col, max_col + 1)
            }
            cpu_rows = [
                row
                for row in range(min_row + 1, max_row + 1)
                if sheet.cell(row, headers["sheet_name"]).value == "CPU"
            ]
            first, second = cpu_rows[:2]
            duplicate_path = sheet.cell(first, headers["json_path"]).value
            sheet.cell(second, headers["json_path"]).value = duplicate_path
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertTrue(
                any("JSONパスが重複" in error for error in result.errors)
            )
        finally:
            temporary.cleanup()

    def test_related_sheets_are_merged_into_product(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            fill_minimum_ssd(workbook)
            write_first_empty_row(
                workbook["Identifiers"],
                "IdentifiersCatalog",
                {
                    "product_id": "test-ssd",
                    "identifier_type": "asins",
                    "identifier_value": "B0TEST0001",
                    "region": "JP",
                    "condition": "new",
                    "is_primary": False,
                },
            )
            write_first_empty_row(
                workbook["Evidence"],
                "EvidenceCatalog",
                {
                    "evidence_id": "ev-test-1",
                    "product_id": "test-ssd",
                    "kind": "manufacturer",
                    "title": "Specification",
                    "url": "https://example.com/spec",
                    "checked_at": "2026-07-15",
                    "supports": "specs.capacity_gb\nspecs.nand_type",
                },
            )
            write_first_empty_row(
                workbook["Risks"],
                "RisksCatalog",
                {
                    "risk_id": "risk-test-1",
                    "product_id": "test-ssd",
                    "risk_flag": "component_swap_unverified",
                    "severity": "medium",
                    "status": "monitoring",
                    "summary": "部品変更を継続確認",
                    "evidence_id": "ev-test-1",
                },
            )
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertEqual(result.errors, [])
            product = next(
                product
                for product in result.catalog["products"]
                if product["id"] == "test-ssd"
            )
            self.assertEqual(product["identifiers"]["asins"], ["B0TEST0001"])
            self.assertEqual(product["quality"]["evidence"][0]["id"], "ev-test-1")
            self.assertEqual(
                product["quality"]["risk_flags"], ["component_swap_unverified"]
            )
            self.assertEqual(product["quality"]["risks"][0]["id"], "risk-test-1")
        finally:
            temporary.cleanup()

    def test_unknown_evidence_kind_is_rejected(self):
        temporary, path = self.copy_workbook()
        try:
            workbook = load_workbook(path)
            workbook["Evidence"]["C5"] = "price_blog"
            workbook.save(path)

            result = build_catalog_from_workbook(path)
            self.assertTrue(
                any(
                    "未対応の根拠種別" in error and "price_blog" in error
                    for error in result.errors
                )
            )
        finally:
            temporary.cleanup()

    def test_cell_types_are_strict_and_human_friendly(self):
        self.assertEqual(convert_value("A\nB;C", "string_list"), ["A", "B", "C"])
        self.assertTrue(convert_value("TRUE", "boolean"))
        self.assertEqual(convert_value("2.0", "integer"), 2)
        self.assertEqual(convert_value("2026-07-15", "date"), "2026-07-15")
        with self.assertRaises(ValueError):
            convert_value("2.5", "integer")


if __name__ == "__main__":
    unittest.main()
