#!/usr/bin/env python3
"""Convert the human-maintained quality catalog workbook to runtime JSON.

Category columns are mapped through the FieldDefinitions sheet.  Adding a
specification therefore requires an Excel column and one mapping row, not a
Python code change.
"""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

try:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
except ImportError as exc:  # pragma: no cover - exercised only before setup
    raise SystemExit(
        "openpyxl がありません。python3 -m pip install -r requirements.txt を実行してください。"
    ) from exc

from quality_catalog import validate_catalog


DEFAULT_WORKBOOK = Path(__file__).with_name("catalog") / "quality_catalog.xlsx"
DEFAULT_OUTPUT = Path(__file__).with_name("catalog") / "quality_catalog.json"
SUPPORTED_DATA_TYPES = {
    "string",
    "integer",
    "number",
    "boolean",
    "date",
    "string_list",
    "json",
}
PATH_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9_]*(?:\.[A-Za-z][A-Za-z0-9_]*)*$")
MOTHERBOARD_SLOT_TYPES = {"pcie_expansion", "m2_storage"}
MOTHERBOARD_SLOT_CONNECTIONS = {"CPU", "Chipset", "CPU/Chipset"}
MOTHERBOARD_USB_LOCATIONS = {"rear", "front_header"}
MOTHERBOARD_USB_CONNECTORS = {"Type-A", "Type-C"}


class ExcelCatalogError(ValueError):
    """The workbook structure cannot be converted safely."""


@dataclass
class BuildResult:
    catalog: dict[str, Any]
    errors: list[str]
    warnings: list[str]


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def clean_string(value: Any) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def split_lines(value: Any) -> list[str]:
    if is_blank(value):
        return []
    if isinstance(value, (list, tuple)):
        values: Iterable[Any] = value
    else:
        values = re.split(r"[\r\n;]+", clean_string(value))
    result: list[str] = []
    for item in values:
        text = clean_string(item)
        if text and text not in result:
            result.append(text)
    return result


def parse_boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value in (0, 1):
            return bool(value)
    normalized = clean_string(value).casefold()
    if normalized in {"true", "yes", "y", "1", "有", "あり"}:
        return True
    if normalized in {"false", "no", "n", "0", "無", "なし"}:
        return False
    raise ValueError("TRUE/FALSE のいずれかを入力してください")


def convert_value(value: Any, data_type: str) -> Any:
    """Convert one Excel cell according to a FieldDefinitions data type."""
    if is_blank(value):
        return None
    if data_type == "string":
        return clean_string(value)
    if data_type == "string_list":
        return split_lines(value)
    if data_type == "boolean":
        return parse_boolean(value)
    if data_type == "integer":
        if isinstance(value, bool):
            raise ValueError("整数を入力してください")
        number = float(value)
        if not math.isfinite(number) or not number.is_integer():
            raise ValueError("整数を入力してください")
        return int(number)
    if data_type == "number":
        if isinstance(value, bool):
            raise ValueError("数値を入力してください")
        number = float(value)
        if not math.isfinite(number):
            raise ValueError("有限の数値を入力してください")
        return int(number) if number.is_integer() else number
    if data_type == "date":
        if isinstance(value, dt.datetime):
            value = value.date()
        if isinstance(value, dt.date):
            return value.isoformat()
        try:
            return dt.date.fromisoformat(clean_string(value)).isoformat()
        except ValueError as exc:
            raise ValueError("YYYY-MM-DD形式の日付を入力してください") from exc
    if data_type == "json":
        if isinstance(value, str):
            try:
                return json.loads(value)
            except json.JSONDecodeError as exc:
                raise ValueError(f"JSONが不正です: {exc.msg}") from exc
        return value
    raise ValueError(f"未対応のdata_typeです: {data_type}")


def set_path(target: dict[str, Any], path: str, value: Any) -> None:
    current = target
    parts = path.split(".")
    for part in parts[:-1]:
        existing = current.get(part)
        if existing is None:
            existing = {}
            current[part] = existing
        if not isinstance(existing, dict):
            raise ValueError(f"{part} は既に値として使われているため配下へ書き込めません")
        current = existing
    current[parts[-1]] = value


def table_records(workbook: Any, sheet_name: str, table_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if sheet_name not in workbook.sheetnames:
        raise ExcelCatalogError(f"必須シートがありません: {sheet_name}")
    sheet = workbook[sheet_name]
    if table_name not in sheet.tables:
        raise ExcelCatalogError(f"{sheet_name} にテーブル {table_name} がありません")
    min_col, min_row, max_col, max_row = range_boundaries(sheet.tables[table_name].ref)
    headers = [clean_string(sheet.cell(min_row, column).value) for column in range(min_col, max_col + 1)]
    if any(not header for header in headers):
        raise ExcelCatalogError(f"{sheet_name}.{table_name} に空の列名があります")
    if len(set(headers)) != len(headers):
        raise ExcelCatalogError(f"{sheet_name}.{table_name} に重複した列名があります")
    records: list[dict[str, Any]] = []
    for row_number in range(min_row + 1, max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(min_col, max_col + 1)]
        if all(is_blank(value) for value in values):
            continue
        record = dict(zip(headers, values))
        record["_excel_row"] = row_number
        records.append(record)
    return headers, records


def require_columns(sheet_name: str, headers: list[str], required: Iterable[str]) -> None:
    missing = [column for column in required if column not in headers]
    if missing:
        raise ExcelCatalogError(f"{sheet_name} に必須列がありません: {', '.join(missing)}")


def row_label(sheet_name: str, row: dict[str, Any]) -> str:
    return f"{sheet_name}!{row.get('_excel_row', '?')}行"


def parse_json_cell(value: Any, location: str, errors: list[str]) -> Any:
    try:
        return convert_value(value, "json")
    except (TypeError, ValueError) as exc:
        errors.append(f"{location}: {exc}")
        return None


def append_unique(values: list[str], value: str) -> None:
    if value not in values:
        values.append(value)


def product_defaults(product: dict[str, Any]) -> None:
    product.setdefault("aliases", [])
    product.setdefault("identifiers", {})
    product.setdefault("specs", {})
    product.setdefault("notes", [])
    quality = product.setdefault("quality", {})
    quality.setdefault("summary", "")
    quality.setdefault("strengths", [])
    quality.setdefault("weaknesses", [])
    quality.setdefault("risk_flags", [])
    quality.setdefault("reviewed_at", None)
    quality.setdefault("evidence", [])


def attach_motherboard_children(
    workbook: Any,
    products_by_id: dict[str, dict[str, Any]],
    errors: list[str],
) -> None:
    """Attach normalized motherboard slot and USB rows to their parent products."""
    slot_headers, slot_rows = table_records(
        workbook, "MotherboardSlots", "MotherboardSlotsCatalog"
    )
    require_columns(
        "MotherboardSlots",
        slot_headers,
        [
            "slot_record_id",
            "product_id",
            "slot_name",
            "slot_type",
            "interface_generation",
            "lane_width",
            "connected_to",
            "supported_sizes",
            "heatsink",
            "shared_with",
            "sharing_effect",
            "availability_condition",
            "notes",
        ],
    )
    slot_ids: set[str] = set()
    for row in slot_rows:
        location = row_label("MotherboardSlots", row)
        slot_id = clean_string(row.get("slot_record_id"))
        product_id = clean_string(row.get("product_id"))
        slot_name = clean_string(row.get("slot_name"))
        slot_type = clean_string(row.get("slot_type"))
        if not slot_id or not product_id or not slot_name or not slot_type:
            errors.append(
                f"{location}: slot_record_id・product_id・slot_name・slot_typeは必須です"
            )
            continue
        if slot_type not in MOTHERBOARD_SLOT_TYPES:
            errors.append(
                f"{location}.slot_type: 未対応値です: {slot_type} "
                f"({', '.join(sorted(MOTHERBOARD_SLOT_TYPES))})"
            )
            continue
        if slot_id in slot_ids:
            errors.append(f"Motherboard slot IDが重複しています: {slot_id}")
            continue
        slot_ids.add(slot_id)
        product = products_by_id.get(product_id)
        if product is None:
            errors.append(f"{location}: 未登録のproduct_idです: {product_id}")
            continue
        if product.get("category") != "motherboard":
            errors.append(f"{location}: motherboardカテゴリのproduct_idではありません")
            continue
        slot: dict[str, Any] = {
            "id": slot_id,
            "name": slot_name,
            "type": slot_type,
        }
        for source, target, data_type in (
            ("interface_generation", "interface_generation", "integer"),
            ("lane_width", "lane_width", "integer"),
            ("connected_to", "connected_to", "string"),
            ("supported_sizes", "supported_sizes", "string_list"),
            ("heatsink", "heatsink", "boolean"),
            ("shared_with", "shared_with", "string_list"),
            ("sharing_effect", "sharing_effect", "string"),
            ("availability_condition", "availability_condition", "string"),
            ("notes", "notes", "string_list"),
        ):
            if is_blank(row.get(source)):
                continue
            try:
                value = convert_value(row[source], data_type)
            except (TypeError, ValueError) as exc:
                errors.append(f"{location}.{source}: {exc}")
                continue
            if target in {"interface_generation", "lane_width"} and value <= 0:
                errors.append(f"{location}.{source}: 1以上の整数を入力してください")
                continue
            if target == "connected_to" and value not in MOTHERBOARD_SLOT_CONNECTIONS:
                errors.append(
                    f"{location}.{source}: 未対応値です: {value} "
                    f"({', '.join(sorted(MOTHERBOARD_SLOT_CONNECTIONS))})"
                )
                continue
            slot[target] = value
        product["specs"].setdefault("slots", []).append(slot)

    usb_headers, usb_rows = table_records(
        workbook, "MotherboardUSB", "MotherboardUSBCatalog"
    )
    require_columns(
        "MotherboardUSB",
        usb_headers,
        [
            "usb_record_id",
            "product_id",
            "location",
            "official_standard",
            "usb_max_speed_gbps",
            "connector_type",
            "port_count",
            "alternate_protocols",
            "features",
            "notes",
        ],
    )
    usb_ids: set[str] = set()
    for row in usb_rows:
        location = row_label("MotherboardUSB", row)
        usb_id = clean_string(row.get("usb_record_id"))
        product_id = clean_string(row.get("product_id"))
        port_location = clean_string(row.get("location"))
        standard = clean_string(row.get("official_standard"))
        connector_type = clean_string(row.get("connector_type"))
        if (
            not usb_id
            or not product_id
            or not port_location
            or not standard
            or not connector_type
        ):
            errors.append(
                f"{location}: usb_record_id・product_id・location・official_standard・"
                "connector_typeは必須です"
            )
            continue
        if port_location not in MOTHERBOARD_USB_LOCATIONS:
            errors.append(
                f"{location}.location: 未対応値です: {port_location} "
                f"({', '.join(sorted(MOTHERBOARD_USB_LOCATIONS))})"
            )
            continue
        if connector_type not in MOTHERBOARD_USB_CONNECTORS:
            errors.append(
                f"{location}.connector_type: 未対応値です: {connector_type} "
                f"({', '.join(sorted(MOTHERBOARD_USB_CONNECTORS))})"
            )
            continue
        if usb_id in usb_ids:
            errors.append(f"Motherboard USB IDが重複しています: {usb_id}")
            continue
        usb_ids.add(usb_id)
        product = products_by_id.get(product_id)
        if product is None:
            errors.append(f"{location}: 未登録のproduct_idです: {product_id}")
            continue
        if product.get("category") != "motherboard":
            errors.append(f"{location}: motherboardカテゴリのproduct_idではありません")
            continue
        usb: dict[str, Any] = {
            "id": usb_id,
            "location": port_location,
            "official_standard": standard,
            "connector_type": connector_type,
        }
        for source, target, data_type in (
            ("usb_max_speed_gbps", "usb_max_speed_gbps", "number"),
            ("port_count", "port_count", "integer"),
            ("alternate_protocols", "alternate_protocols", "string_list"),
            ("features", "features", "string_list"),
            ("notes", "notes", "string_list"),
        ):
            if is_blank(row.get(source)):
                if source == "port_count":
                    errors.append(f"{location}.port_count は必須です")
                continue
            try:
                value = convert_value(row[source], data_type)
            except (TypeError, ValueError) as exc:
                errors.append(f"{location}.{source}: {exc}")
                continue
            if target == "port_count" and value <= 0:
                errors.append(f"{location}.port_count: 1以上の整数を入力してください")
                continue
            if target == "usb_max_speed_gbps" and value <= 0:
                errors.append(
                    f"{location}.usb_max_speed_gbps: 0より大きい数値を入力してください"
                )
                continue
            usb[target] = value
        product["specs"].setdefault("usb_ports", []).append(usb)


def validate_motherboard_usb_aggregates(
    products: list[dict[str, Any]], errors: list[str]
) -> None:
    """Verify parent-sheet USB totals against normalized rear-I/O rows."""
    for product in products:
        if product.get("category") != "motherboard":
            continue
        specs = product.get("specs") or {}
        rear = [
            row
            for row in specs.get("usb_ports") or []
            if clean_string(row.get("location")).casefold() == "rear"
        ]
        if not rear:
            continue
        bad_connectors = sorted(
            {
                clean_string(row.get("connector_type"))
                for row in rear
                if clean_string(row.get("connector_type")) not in {"Type-A", "Type-C"}
            }
        )
        if bad_connectors:
            errors.append(
                f"{product['id']}.specs.usb_ports: 背面USBのconnector_typeはType-A/Type-Cのみです: "
                + ", ".join(bad_connectors)
            )
            continue
        if any(not isinstance(row.get("port_count"), int) for row in rear):
            errors.append(f"{product['id']}.specs.usb_ports: 背面USBのport_countが未入力です")
            continue
        expected = {
            "rear_usb_total_count": sum(row["port_count"] for row in rear),
            "rear_usb_type_a_count": sum(
                row["port_count"] for row in rear if row["connector_type"] == "Type-A"
            ),
            "rear_usb_type_c_count": sum(
                row["port_count"] for row in rear if row["connector_type"] == "Type-C"
            ),
            "rear_usb4_type_c_count": sum(
                row["port_count"]
                for row in rear
                if row["connector_type"] == "Type-C"
                and "usb4" in clean_string(row.get("official_standard")).casefold()
            ),
        }
        speeds = [
            row.get("usb_max_speed_gbps")
            for row in rear
            if isinstance(row.get("usb_max_speed_gbps"), (int, float))
            and not isinstance(row.get("usb_max_speed_gbps"), bool)
        ]
        if len(speeds) != len(rear):
            errors.append(
                f"{product['id']}.specs.usb_ports: 背面USBのusb_max_speed_gbpsが未入力です"
            )
        else:
            expected["rear_usb_fastest_gbps"] = max(speeds)
        for key, expected_value in expected.items():
            actual = specs.get(key)
            if actual != expected_value:
                errors.append(
                    f"{product['id']}.specs.{key}: MotherboardUSB集計は"
                    f"{expected_value}ですがMotherboard値は{actual!r}です"
                )


def build_catalog_from_workbook(path: Path) -> BuildResult:
    """Build a catalog and return all actionable workbook errors/warnings."""
    errors: list[str] = []
    warnings: list[str] = []
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except (OSError, ValueError) as exc:
        raise ExcelCatalogError(f"Excelファイルを読み込めません: {path}: {exc}") from exc

    sheet_headers, sheet_rows = table_records(
        workbook, "SheetDefinitions", "SheetDefinitionsCatalog"
    )
    require_columns(
        "SheetDefinitions",
        sheet_headers,
        ["sheet_name", "category", "table_name", "enabled"],
    )
    definitions: list[dict[str, str]] = []
    for row in sheet_rows:
        location = row_label("SheetDefinitions", row)
        try:
            enabled = parse_boolean(row.get("enabled"))
        except ValueError as exc:
            errors.append(f"{location}.enabled: {exc}")
            continue
        if not enabled:
            continue
        values = {key: clean_string(row.get(key)) for key in ("sheet_name", "category", "table_name")}
        for key, value in values.items():
            if not value:
                errors.append(f"{location}.{key} は必須です")
        if all(values.values()):
            definitions.append(values)

    field_headers, field_rows = table_records(
        workbook, "FieldDefinitions", "FieldDefinitionsCatalog"
    )
    require_columns(
        "FieldDefinitions",
        field_headers,
        ["sheet_name", "column_name", "json_path", "data_type", "required", "active"],
    )
    field_definitions: dict[str, list[dict[str, Any]]] = {}
    field_keys: set[tuple[str, str]] = set()
    field_paths: set[tuple[str, str]] = set()
    for row in field_rows:
        location = row_label("FieldDefinitions", row)
        try:
            active = parse_boolean(row.get("active"))
        except ValueError as exc:
            errors.append(f"{location}.active: {exc}")
            continue
        if not active:
            continue
        sheet_name = clean_string(row.get("sheet_name"))
        column_name = clean_string(row.get("column_name"))
        json_path = clean_string(row.get("json_path"))
        data_type = clean_string(row.get("data_type"))
        try:
            required = parse_boolean(row.get("required"))
        except ValueError as exc:
            errors.append(f"{location}.required: {exc}")
            required = False
        for name, value in (
            ("sheet_name", sheet_name),
            ("column_name", column_name),
            ("json_path", json_path),
            ("data_type", data_type),
        ):
            if not value:
                errors.append(f"{location}.{name} は必須です")
        if data_type and data_type not in SUPPORTED_DATA_TYPES:
            errors.append(f"{location}.data_type={data_type!r} は未対応です")
        if json_path and not PATH_PATTERN.fullmatch(json_path):
            errors.append(f"{location}.json_path={json_path!r} は不正です")
        key = (sheet_name, column_name)
        if all(key):
            if key in field_keys:
                errors.append(f"FieldDefinitionsの定義が重複しています: {sheet_name}.{column_name}")
            field_keys.add(key)
        path_key = (sheet_name, json_path)
        if all(path_key):
            if path_key in field_paths:
                errors.append(
                    f"FieldDefinitionsのJSONパスが重複しています: {sheet_name}.{json_path}"
                )
            field_paths.add(path_key)
        if sheet_name and column_name and json_path and data_type in SUPPORTED_DATA_TYPES:
            field_definitions.setdefault(sheet_name, []).append(
                {
                    "column_name": column_name,
                    "json_path": json_path,
                    "data_type": data_type,
                    "required": required,
                }
            )

    products: list[dict[str, Any]] = []
    products_by_id: dict[str, dict[str, Any]] = {}
    known_sheets = {definition["sheet_name"] for definition in definitions}
    unused_definition_sheets = sorted(set(field_definitions) - known_sheets)
    for sheet_name in unused_definition_sheets:
        warnings.append(f"FieldDefinitions.{sheet_name}: 有効なカテゴリシート定義がありません")

    for definition in definitions:
        sheet_name = definition["sheet_name"]
        headers, rows = table_records(workbook, sheet_name, definition["table_name"])
        fields = field_definitions.get(sheet_name, [])
        if not fields:
            errors.append(f"{sheet_name}: 有効なFieldDefinitionsがありません")
            continue
        mapped_columns = {field["column_name"] for field in fields}
        unknown_columns = [header for header in headers if header not in mapped_columns]
        missing_columns = [field["column_name"] for field in fields if field["column_name"] not in headers]
        if unknown_columns:
            errors.append(
                f"{sheet_name}: FieldDefinitions未登録の列があります: {', '.join(unknown_columns)}"
            )
        if missing_columns:
            errors.append(
                f"{sheet_name}: 定義された列がテーブルにありません: {', '.join(missing_columns)}"
            )
        if "product_id" not in headers:
            errors.append(f"{sheet_name}: product_id列がありません")
            continue
        for row in rows:
            location = row_label(sheet_name, row)
            if is_blank(row.get("product_id")):
                errors.append(f"{location}: 入力済みですがproduct_idが空です")
                continue
            product: dict[str, Any] = {"category": definition["category"]}
            for field in fields:
                column_name = field["column_name"]
                if column_name not in row:
                    continue
                raw = row[column_name]
                if is_blank(raw):
                    if field["required"]:
                        errors.append(f"{location}.{column_name} は必須です")
                    continue
                try:
                    value = convert_value(raw, field["data_type"])
                    set_path(product, field["json_path"], value)
                except (TypeError, ValueError) as exc:
                    errors.append(f"{location}.{column_name}: {exc}")
            product_defaults(product)
            product_id = clean_string(product.get("id"))
            if not product_id:
                continue
            if product_id in products_by_id:
                errors.append(f"製品IDが重複しています: {product_id}")
                continue
            products.append(product)
            products_by_id[product_id] = product

    # GPU board rows keep only a stable reference to the normalized chip row in
    # Excel.  The generated JSON receives a read-only snapshot so profiles can
    # compare chip performance without duplicating it in every board SKU row.
    for product in products:
        if product.get("category") != "gpu":
            continue
        specs = product.get("specs") or {}
        chip_id = clean_string(specs.get("gpu_chip_id"))
        if not chip_id:
            continue
        chip = products_by_id.get(chip_id)
        if chip is None:
            errors.append(f"{product['id']}.specs.gpu_chip_id: 未登録のGPUチップIDです: {chip_id}")
            continue
        if chip.get("category") != "gpu_chip":
            errors.append(
                f"{product['id']}.specs.gpu_chip_id: gpu_chipカテゴリではありません: {chip_id}"
            )
            continue
        chip_snapshot = copy.deepcopy(chip.get("specs") or {})
        chip_snapshot.update(
            {
                "id": chip["id"],
                "brand": chip.get("brand"),
                "model": chip.get("model"),
                "display_name": chip.get("display_name"),
            }
        )
        specs["gpu_chip"] = chip_snapshot

    attach_motherboard_children(workbook, products_by_id, errors)
    validate_motherboard_usb_aggregates(products, errors)

    identifier_headers, identifier_rows = table_records(
        workbook, "Identifiers", "IdentifiersCatalog"
    )
    require_columns(
        "Identifiers",
        identifier_headers,
        ["product_id", "identifier_type", "identifier_value"],
    )
    for row in identifier_rows:
        location = row_label("Identifiers", row)
        product_id = clean_string(row.get("product_id"))
        identifier_type = clean_string(row.get("identifier_type"))
        identifier_value = clean_string(row.get("identifier_value"))
        if not product_id or not identifier_type or not identifier_value:
            errors.append(f"{location}: product_id・identifier_type・identifier_valueは必須です")
            continue
        product = products_by_id.get(product_id)
        if product is None:
            errors.append(f"{location}: 未登録のproduct_idです: {product_id}")
            continue
        values = product["identifiers"].setdefault(identifier_type, [])
        append_unique(values, identifier_value)
        detail: dict[str, Any] = {"type": identifier_type, "value": identifier_value}
        for source, target, data_type in (
            ("region", "region", "string"),
            ("condition", "condition", "string"),
            ("is_primary", "is_primary", "boolean"),
            ("notes", "notes", "string"),
        ):
            if not is_blank(row.get(source)):
                try:
                    detail[target] = convert_value(row[source], data_type)
                except ValueError as exc:
                    errors.append(f"{location}.{source}: {exc}")
        product.setdefault("identifier_records", []).append(detail)

    evidence_headers, evidence_rows = table_records(workbook, "Evidence", "EvidenceCatalog")
    require_columns(
        "Evidence",
        evidence_headers,
        ["evidence_id", "product_id", "kind", "title", "url", "checked_at"],
    )
    evidence_owners: dict[str, str] = {}
    for row in evidence_rows:
        location = row_label("Evidence", row)
        required = {
            key: clean_string(row.get(key))
            for key in ("evidence_id", "product_id", "kind", "title", "url")
        }
        if any(not value for value in required.values()) or is_blank(row.get("checked_at")):
            errors.append(f"{location}: evidence_id・product_id・kind・title・url・checked_atは必須です")
            continue
        evidence_id = required["evidence_id"]
        if evidence_id in evidence_owners:
            errors.append(f"Evidence IDが重複しています: {evidence_id}")
            continue
        product = products_by_id.get(required["product_id"])
        if product is None:
            errors.append(f"{location}: 未登録のproduct_idです: {required['product_id']}")
            continue
        parsed_url = urlparse(required["url"])
        if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
            errors.append(f"{location}.url: http/httpsの完全なURLを入力してください")
        try:
            checked_at = convert_value(row.get("checked_at"), "date")
        except ValueError as exc:
            errors.append(f"{location}.checked_at: {exc}")
            continue
        evidence: dict[str, Any] = {
            "id": evidence_id,
            "kind": required["kind"],
            "title": required["title"],
            "url": required["url"],
            "checked_at": checked_at,
        }
        supports = split_lines(row.get("supports"))
        if supports:
            evidence["supports"] = supports
        for key in ("reviewer", "notes"):
            if not is_blank(row.get(key)):
                evidence[key] = clean_string(row[key])
        product["quality"]["evidence"].append(evidence)
        evidence_owners[evidence_id] = required["product_id"]

    risk_headers, risk_rows = table_records(workbook, "Risks", "RisksCatalog")
    require_columns(
        "Risks",
        risk_headers,
        ["risk_id", "product_id", "risk_flag", "severity", "status", "summary"],
    )
    risk_ids: set[str] = set()
    for row in risk_rows:
        location = row_label("Risks", row)
        required = {
            key: clean_string(row.get(key))
            for key in ("risk_id", "product_id", "risk_flag", "severity", "status", "summary")
        }
        if any(not value for value in required.values()):
            errors.append(f"{location}: risk_id・product_id・risk_flag・severity・status・summaryは必須です")
            continue
        risk_id = required["risk_id"]
        if risk_id in risk_ids:
            errors.append(f"Risk IDが重複しています: {risk_id}")
            continue
        product = products_by_id.get(required["product_id"])
        if product is None:
            errors.append(f"{location}: 未登録のproduct_idです: {required['product_id']}")
            continue
        risk: dict[str, Any] = {
            "id": risk_id,
            "flag": required["risk_flag"],
            "severity": required["severity"],
            "status": required["status"],
            "summary": required["summary"],
        }
        for key in ("workaround", "evidence_id", "notes"):
            if not is_blank(row.get(key)):
                risk[key] = clean_string(row[key])
        for key in ("first_checked_at", "last_checked_at"):
            if not is_blank(row.get(key)):
                try:
                    risk[key] = convert_value(row[key], "date")
                except ValueError as exc:
                    errors.append(f"{location}.{key}: {exc}")
        evidence_id = risk.get("evidence_id")
        if evidence_id and evidence_owners.get(evidence_id) != required["product_id"]:
            errors.append(f"{location}.evidence_id: 同じ製品のEvidence IDではありません")
        product["quality"].setdefault("risks", []).append(risk)
        if required["status"] != "resolved":
            append_unique(product["quality"]["risk_flags"], required["risk_flag"])
        risk_ids.add(risk_id)

    settings_headers, settings_rows = table_records(workbook, "Settings", "SettingsCatalog")
    require_columns("Settings", settings_headers, ["key", "value_json"])
    settings: dict[str, Any] = {}
    updated_at: str | None = None
    setting_keys: set[str] = set()
    for row in settings_rows:
        location = row_label("Settings", row)
        key = clean_string(row.get("key"))
        if not key or is_blank(row.get("value_json")):
            errors.append(f"{location}: keyとvalue_jsonは必須です")
            continue
        if key in setting_keys:
            errors.append(f"Settings keyが重複しています: {key}")
            continue
        value = parse_json_cell(row.get("value_json"), f"{location}.value_json", errors)
        setting_keys.add(key)
        if key == "catalog.updated_at":
            try:
                updated_at = convert_value(value, "date")
            except (TypeError, ValueError) as exc:
                errors.append(f"{location}.value_json: {exc}")
            continue
        try:
            set_path(settings, key, value)
        except ValueError as exc:
            errors.append(f"{location}.key: {exc}")
    if updated_at is None:
        errors.append("Settingsにcatalog.updated_atがありません")

    profile_headers, profile_rows = table_records(workbook, "Profiles", "ProfilesCatalog")
    require_columns(
        "Profiles",
        profile_headers,
        ["profile_id", "category", "name", "enabled"],
    )
    profiles: list[dict[str, Any]] = []
    profiles_by_id: dict[str, dict[str, Any]] = {}
    gate_fields = (
        ("minimum_evidence_count", "integer"),
        ("required_evidence_kinds", "string_list"),
        ("forbidden_risk_flags", "string_list"),
        ("minimum_tier", "string"),
        ("max_review_age_days", "integer"),
        ("allowed_statuses", "string_list"),
    )
    for row in profile_rows:
        location = row_label("Profiles", row)
        try:
            enabled = parse_boolean(row.get("enabled"))
        except ValueError as exc:
            errors.append(f"{location}.enabled: {exc}")
            continue
        if not enabled:
            continue
        profile_id = clean_string(row.get("profile_id"))
        category = clean_string(row.get("category"))
        name = clean_string(row.get("name"))
        if not profile_id or not category or not name:
            errors.append(f"{location}: profile_id・category・nameは必須です")
            continue
        if profile_id in profiles_by_id:
            errors.append(f"Profile IDが重複しています: {profile_id}")
            continue
        profile: dict[str, Any] = {
            "id": profile_id,
            "name": name,
            "category": category,
            "requirements": [],
            "preferences": [],
        }
        if not is_blank(row.get("description")):
            profile["description"] = clean_string(row["description"])
        gate: dict[str, Any] = {}
        for key, data_type in gate_fields:
            if is_blank(row.get(key)):
                continue
            try:
                gate[key] = convert_value(row[key], data_type)
            except (TypeError, ValueError) as exc:
                errors.append(f"{location}.{key}: {exc}")
        if gate:
            profile["quality_gate"] = gate
        profiles.append(profile)
        profiles_by_id[profile_id] = profile

    rule_headers, rule_rows = table_records(workbook, "Rules", "RulesCatalog")
    require_columns(
        "Rules",
        rule_headers,
        ["profile_id", "rule_type", "path", "op", "value_json", "enabled"],
    )
    for row in rule_rows:
        location = row_label("Rules", row)
        try:
            enabled = parse_boolean(row.get("enabled"))
        except ValueError as exc:
            errors.append(f"{location}.enabled: {exc}")
            continue
        if not enabled:
            continue
        profile_id = clean_string(row.get("profile_id"))
        rule_type = clean_string(row.get("rule_type"))
        path_value = clean_string(row.get("path"))
        operation = clean_string(row.get("op"))
        if not profile_id or not rule_type or not path_value or not operation or is_blank(row.get("value_json")):
            errors.append(f"{location}: profile_id・rule_type・path・op・value_jsonは必須です")
            continue
        profile = profiles_by_id.get(profile_id)
        if profile is None:
            errors.append(f"{location}: 有効なProfile IDではありません: {profile_id}")
            continue
        if rule_type not in {"requirement", "preference"}:
            errors.append(f"{location}.rule_type={rule_type!r} は未対応です")
            continue
        value = parse_json_cell(row.get("value_json"), f"{location}.value_json", errors)
        rule: dict[str, Any] = {"path": path_value, "op": operation, "value": value}
        if not is_blank(row.get("label")):
            rule["label"] = clean_string(row["label"])
        if rule_type == "preference":
            if is_blank(row.get("weight")):
                errors.append(f"{location}.weight: preferenceでは必須です")
            else:
                try:
                    rule["weight"] = convert_value(row["weight"], "number")
                except (TypeError, ValueError) as exc:
                    errors.append(f"{location}.weight: {exc}")
        profile["requirements" if rule_type == "requirement" else "preferences"].append(rule)

    catalog = {
        "schema_version": 1,
        "updated_at": updated_at,
        "settings": settings,
        "profiles": profiles,
        "products": products,
    }
    validation_errors, validation_warnings = validate_catalog(catalog)
    errors.extend(f"JSON検証: {error}" for error in validation_errors)
    warnings.extend(validation_warnings)
    return BuildResult(catalog=catalog, errors=errors, warnings=warnings)


def print_result(result: BuildResult) -> None:
    for warning in result.warnings:
        print(f"警告: {warning}")
    for error in result.errors:
        print(f"エラー: {error}", file=sys.stderr)


def write_catalog(path: Path, catalog: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="品質カタログExcelの検証・JSON生成")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("validate", help="Excel全体を検証する")
    build_parser = subparsers.add_parser("build", help="検証後にJSONを生成する")
    build_parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    check_parser = subparsers.add_parser("check", help="既存JSONとExcelが一致するか確認する")
    check_parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)

    try:
        result = build_catalog_from_workbook(args.workbook)
    except ExcelCatalogError as exc:
        print(f"エラー: {exc}", file=sys.stderr)
        return 2
    print_result(result)
    if result.errors:
        print(f"検証失敗: {len(result.errors)}件", file=sys.stderr)
        return 1

    if args.command == "build":
        write_catalog(args.output, result.catalog)
        print(
            f"JSON生成完了: {args.output} "
            f"(products={len(result.catalog['products'])}, profiles={len(result.catalog['profiles'])})"
        )
    elif args.command == "check":
        try:
            current = json.loads(args.output.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            print(f"エラー: 比較対象JSONを読み込めません: {args.output}: {exc}", file=sys.stderr)
            return 2
        if current != result.catalog:
            print(
                "エラー: ExcelとJSONが一致しません。catalog_excel.py buildを実行してください。",
                file=sys.stderr,
            )
            return 1
        print(f"同期確認完了: {args.workbook} == {args.output}")
    else:
        print(
            f"検証完了: products={len(result.catalog['products'])}, "
            f"profiles={len(result.catalog['profiles'])}, warnings={len(result.warnings)}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
